#!/usr/bin/env python3
"""Turn an .xlsx export of the "Μηνιαία Έξοδα" Google Sheet into SQL for Supabase.

    python3 supabase/import_from_xlsx.py export.xlsx you@gmail.com partner@gmail.com > data.sql

Then paste data.sql into the Supabase SQL Editor (after schema.sql) and run it.
It creates one household, lets the given emails in, and copies payers, categories,
monthly limits, plans and every payment. The output contains your budget data:
keep it private and never commit it to the repository.

Needs: pip install openpyxl
"""
import datetime
import sys

import openpyxl

LOG, PLANS, DASH = 'Καταγραφή', 'Ιστορικό και πλάνα', 'Έλεγχος μήνα'
DEFAULT_PAYERS = ['Αθηνά', 'Γιάννης', 'Κοινό']


def q(v):
    """SQL literal."""
    if v is None:
        return 'null'
    if isinstance(v, (int, float)):
        return repr(round(float(v), 2))
    return "'" + str(v).replace("'", "''") + "'"


def month_of(d):
    return d.strftime('%Y-%m-01')


def main(path, emails):
    wb = openpyxl.load_workbook(path, data_only=True)
    log, plans, dash = wb[LOG], wb[PLANS], wb[DASH]

    categories = [str(r[0].value).strip() for r in dash['A16:A32'] if r[0].value not in (None, '')]

    payers = DEFAULT_PAYERS
    for dv in log.data_validations.dataValidation:
        if 'E5' in str(dv.sqref) and dv.formula1 and dv.formula1.startswith('"'):
            payers = [p.strip() for p in dv.formula1.strip('"').split(',') if p.strip()]

    months, plan_rows, default_type = [], [], {}
    for r in plans.iter_rows(min_row=5, values_only=True):
        if isinstance(r[0], datetime.datetime):
            months.append((r[0], r[1] or 0, r[2] if r[2] not in (None, '') else 0.8))
        if isinstance(r[7], datetime.datetime) and r[8]:
            kind = 'fixed' if r[9] == 'Πάγιο' else 'extra'
            plan_rows.append((r[7], str(r[8]), kind, r[10] if r[10] not in (None, '') else None))
            default_type.setdefault(str(r[8]), kind)

    payments = []
    for r in log.iter_rows(min_row=5, max_row=2001, max_col=6, values_only=True):
        date, cat, desc, amount, payer, comments = r
        if date in (None, '') and cat in (None, '') and amount in (None, ''):
            continue
        if not isinstance(date, (datetime.datetime, datetime.date)) or not cat or not amount:
            print(f'-- skipped incomplete row: {r!r}', file=sys.stderr)
            continue
        if cat not in categories:
            categories.append(str(cat))
        payments.append((date, str(cat), desc or '', float(amount), payer or '', comments or ''))

    out = ['begin;', '',
           "do $$ begin if exists (select 1 from payments where household_id = '00000000-0000-4000-8000-000000000001') "
           "then raise exception 'Payments were already imported. Nothing was changed.'; end if; end $$;", '',
           "insert into households (id, name) values ('00000000-0000-4000-8000-000000000001', 'Οικογένεια')",
           '  on conflict do nothing;',
           "create temporary table _h on commit drop as select '00000000-0000-4000-8000-000000000001'::uuid as id;", '']
    for e in emails:
        out.append(f'insert into allowed_emails (email, household_id) select {q(e.lower())}, id from _h on conflict (email) do nothing;')
    out.append('')
    for i, p in enumerate(payers):
        out.append(f'insert into payers (household_id, name, sort_order) select id, {q(p)}, {i} from _h on conflict do nothing;')
    out.append('')
    for i, c in enumerate(categories):
        out.append(f"insert into categories (household_id, name, default_type, sort_order) "
                   f"select id, {q(c)}, '{default_type.get(c, 'extra')}', {i} from _h on conflict do nothing;")
    out.append('')
    for m, cap, warn in months:
        out.append(f'insert into month_limits (household_id, month, cap, warn_ratio) '
                   f'select id, {q(month_of(m))}, {q(cap)}, {q(warn)} from _h on conflict do nothing;')
    out.append('')
    for m, cat, kind, planned in plan_rows:
        out.append(f"insert into category_plans (household_id, month, category_id, type, planned) "
                   f"select h.id, {q(month_of(m))}, c.id, '{kind}', {q(planned)} from _h h "
                   f"join categories c on c.household_id = h.id and c.name = {q(cat)} on conflict do nothing;")
    out.append('')
    for date, cat, desc, amount, payer, comments in payments:
        out.append(f"insert into payments (household_id, paid_on, category_id, payer_id, amount, description, comments) "
                   f"select h.id, {q(date.strftime('%Y-%m-%d'))}, c.id, "
                   f"(select y.id from payers y where y.household_id = h.id and y.name = {q(payer)}), "
                   f"{q(amount)}, {q(str(desc))}, {q(str(comments))} from _h h "
                   f"join categories c on c.household_id = h.id and c.name = {q(cat)};")
    out += ['', 'commit;', '',
            f'-- {len(payers)} payers, {len(categories)} categories, {len(months)} months, '
            f'{len(plan_rows)} plan rows, {len(payments)} payments']
    print('\n'.join(out))


if __name__ == '__main__':
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2:])
